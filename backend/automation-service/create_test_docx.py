#!/usr/bin/env python3
"""Create a test DOCX file with contract content"""

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create a new Document
doc = Document()

# Add title
title = doc.add_heading('HỢP ĐỒNG KINH DOANH', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Add content
doc.add_paragraph('Hợp đồng này được ký kết giữa:')
doc.add_paragraph('Bên A: Công ty TNHH ABC', style='List Bullet')
doc.add_paragraph('Bên B: Công ty TNHH XYZ', style='List Bullet')

doc.add_heading('I. Mục đích hợp tác', level=1)
doc.add_paragraph('Hai bên thỏa thuận hợp tác kinh doanh với mục đích phát triển sản phẩm và dịch vụ.')

doc.add_heading('II. Điều khoản và điều kiện', level=1)
doc.add_paragraph('1. Bên A cam kết cung cấp dịch vụ theo đúng tiêu chuẩn.')
doc.add_paragraph('2. Bên B cam kết thanh toán đúng hạn theo hóa đơn.')
doc.add_paragraph('3. Thời hạn hợp đồng: 2 năm kể từ ngày ký kết.')

doc.add_heading('III. Giá trị hợp đồng', level=1)
doc.add_paragraph('Tổng giá trị: 100,000,000 VND')
doc.add_paragraph('Phương thức thanh toán: Chuyển khoản ngân hàng')

doc.add_heading('IV. Ký kết', level=1)
doc.add_paragraph('Hợp đồng này được ký kết tại Hà Nội, ngày 25 tháng 10 năm 2025')
doc.add_paragraph('Bên A: ________________')
doc.add_paragraph('Bên B: ________________')

# Save the document
output_path = 'P:\\DevGO2003\\DocGO-private-new\\.cursor\\documents\\.docx\\luu-ban-nhap-tu-dong-2.docx'
doc.save(output_path)
print(f"✓ Test DOCX file created: {output_path}")
